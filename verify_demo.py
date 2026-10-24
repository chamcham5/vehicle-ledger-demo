"""Reproduce the fictional demonstration using only this repository.

Temporary workbooks, CSVs and changed source inputs are removed afterwards.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import subprocess
import sys
import tempfile
from datetime import date, datetime
from pathlib import Path

sys.dont_write_bytecode = True
from openpyxl import load_workbook
import pipeline

EXPECTED_JANUARY_CSV_SHA256 = "d1f3f35bb331c4b425015d2e2f5dccaadd59887a426eb1509b093d98c0131c68"

HERE = Path(__file__).resolve().parent
BASE = HERE


def invoke(vehicles, out):
    return subprocess.run([sys.executable, "-B", str(HERE / "pipeline.py"),
                           "--vehicles", str(vehicles),
                           "--contracts", str(BASE / "contracts.pdf"),
                           "--installments", str(BASE / "installments.pdf"),
                           "--month", "2027-01", "--out", str(out)],
                          capture_output=True, text=True, check=False)


def verify():
    checks = {}
    with tempfile.TemporaryDirectory(prefix="vehicle-ledger-demo-") as folder:
        temp = Path(folder).resolve()
        assert temp.parent == Path(tempfile.gettempdir()).resolve()
        assert temp.name.startswith("vehicle-ledger-demo-")
        destination = temp / "new-result"
        normal = invoke(BASE / "vehicles.xlsx", destination)
        assert normal.returncode == 0, normal.stderr
        cli = json.loads(normal.stdout)
        checks["actual_cli_three_inputs_to_new_saved_outputs"] = cli["payment_rows"] == 6
        assert sorted(p.name for p in destination.iterdir()) == [
            "finance.csv", "payment-schedule.xlsx", "source-report.json"]
        report = json.loads((destination / "source-report.json").read_text(encoding="utf-8"))
        checks["report_generated_from_only_three_fixture_inputs"] = (
            [v["name"] for v in report["source_inputs"]] ==
            ["vehicles.xlsx", "contracts.pdf", "installments.pdf"] and
            report["pdf_pages"] == {"contracts": 3, "installments": 3} and
            len(report["vehicles"]) == 3)
        saved = (destination / "payment-schedule.xlsx").read_bytes()
        assert hashlib.sha256(saved).hexdigest() == cli["saved_xlsx_sha256"]
        book = load_workbook(io.BytesIO(saved), read_only=True, data_only=True)
        try:
            details = list(book["Payments"].iter_rows(min_row=5, max_row=10, values_only=True))
            checks["generated_saved_xlsx_typed_rows"] = (
                book.sheetnames == ["Payments", "Reconciliation", "Exceptions"] and
                all(isinstance(r[3], (datetime, date)) and type(r[4]) is int for r in details) and
                sum(r[4] for r in details) == 545000)
        finally:
            book.close()
        # The expected fixture hash is comparison-only, after generation has completed.
        checks["generated_january_csv_byte_matches_golden"] = (
            hashlib.sha256((destination / "finance.csv").read_bytes()).hexdigest() == EXPECTED_JANUARY_CSV_SHA256)
        months = {}
        for month in ("2027-01", "2027-02", "2027-03", "2027-04"):
            data = pipeline.WRITER.export_csv(saved, report, month)
            rows = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
            months[month] = {"rows": len(rows), "total_jpy": sum(int(r["amount_jpy"]) for r in rows)}
        checks["four_months_from_same_new_saved_xlsx"] = (
            [m["total_jpy"] for m in months.values()] == [225001, 160000, 159999, 0])

        bad_source = temp / "vehicles-mismatch.xlsx"
        book = load_workbook(BASE / "vehicles.xlsx")
        original_registration = book["Vehicles"]["B2"].value
        book["Vehicles"]["B2"] = original_registration + "_MISMATCH"
        book.save(bad_source)
        book.close()
        blocked_path = temp / "must-not-exist"
        blocked = invoke(bad_source, blocked_path)
        blocked_result = json.loads(blocked.stderr)
        checks["actual_source_mismatch_blocks_before_output"] = (
            blocked.returncode == 2 and not blocked_path.exists() and
            blocked_result == {"status": "BLOCKED", "error": "Vehicle/contract mismatch"})

        book = load_workbook(bad_source)
        book["Vehicles"]["B2"] = original_registration
        book.save(bad_source)
        book.close()
        recovered = invoke(bad_source, blocked_path)
        checks["manual_source_correction_then_new_run_recovers"] = (
            recovered.returncode == 0 and
            hashlib.sha256((blocked_path / "finance.csv").read_bytes()).hexdigest() == EXPECTED_JANUARY_CSV_SHA256)
        assert all(checks.values()), checks
    assert not Path(folder).exists()
    result = {"status": "PASS", "scope": "Self-contained fictional repository example only",
              "checks": checks, "passed": len(checks), "monthly_results": months,
              "normal_cli": cli, "blocked_cli": blocked_result,
              "temporary_workbooks_csv_and_modified_input_removed": True,
              "real_client_format_or_gui_verified": False,
              "requires_private_reference_files": False}
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    verify()
