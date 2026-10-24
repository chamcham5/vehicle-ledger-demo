"""Fixed fictional-schema demonstration, not a client-ready parser.

Inputs are the vehicle XLSX and two tagged text PDFs only. Existing demo
outputs are never generation inputs. The saved new XLSX supplies CSV rows.
"""
from __future__ import annotations

import argparse
import hashlib
import importlib.util
import json
import re
import sys
from datetime import date
from pathlib import Path

sys.dont_write_bytecode = True
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent


import source_guards as READER
import workbook_writer as WRITER

def extract_report(vehicles_path, contracts_path, installments_path, month):
    month = WRITER.month_key(month)
    inputs = [Path(p).resolve(strict=True) for p in
              (vehicles_path, contracts_path, installments_path)]
    book = load_workbook(inputs[0], read_only=True, data_only=False)
    try:
        rows = list(book["Vehicles"].iter_rows(values_only=True))
    finally:
        book.close()
    if not rows or tuple(rows[0]) != ("vehicle_id", "registration", "model"):
        raise ValueError("Unexpected vehicle header")
    vehicles = {}
    for number, row in enumerate(rows[1:], 2):
        if len(row) != 3 or any(not isinstance(v, str) or not v.strip() for v in row):
            raise ValueError("Missing/invalid vehicle row")
        vehicle, registration, model = row
        if vehicle in vehicles:
            raise ValueError("Duplicate vehicle ID")
        vehicles[vehicle] = {"registration": registration, "model": model,
                             "xlsx_source": f"{inputs[0].name}:Vehicles!{number}"}
    if not vehicles:
        raise ValueError("No vehicles")

    contracts = {}
    contract_pages = READER.read_pdf_pages(inputs[1], "Demo schema: CONTRACT_V1")
    for page, text in contract_pages:
        vehicle = READER.field(text, "Vehicle ID")
        if vehicle not in vehicles or vehicle in contracts:
            raise ValueError("Duplicate/unknown contract vehicle")
        expected = vehicles[vehicle]
        if (READER.field(text, "Registration") != expected["registration"] or
                READER.field(text, "Model") != expected["model"]):
            raise ValueError("Vehicle/contract mismatch")
        contracts[vehicle] = {"total": READER.jint(READER.field(text, "Finance total JPY")),
                              "source": f"{inputs[1].name}#page={page}"}
    if set(contracts) != set(vehicles):
        raise ValueError("Missing vehicle contract")

    payments, seen = [], set()
    schedule_pages = READER.read_pdf_pages(inputs[2], "Demo schema: SCHEDULE_V1")
    for page, text in schedule_pages:
        vehicle = READER.field(text, "Vehicle ID")
        if vehicle not in contracts or vehicle in seen:
            raise ValueError("Duplicate/unknown installment schedule")
        seen.add(vehicle)
        count = READER.jint(READER.field(text, "Installments"))
        total = READER.jint(READER.field(text, "Total JPY"))
        current = []
        for line in re.findall(r"^Payment: ([^\n]+)$", text, re.M):
            parts = line.strip().split("|")
            if len(parts) != 3:
                raise ValueError("Invalid payment line")
            installment, due, amount = parts
            current.append({"vehicle_id": vehicle,
                            "registration": vehicles[vehicle]["registration"],
                            "installment": READER.jint(installment),
                            "payment_date": date.fromisoformat(due).isoformat(),
                            "amount_jpy": READER.jint(amount),
                            "source": f"{inputs[2].name}#page={page}"})
        if count < 1 or [p["installment"] for p in current] != list(range(1, count + 1)):
            raise ValueError("Invalid installment sequence/count")
        if sum(p["amount_jpy"] for p in current) != total or total != contracts[vehicle]["total"]:
            raise ValueError("Schedule/contract total mismatch")
        payments.extend(current)
    if seen != set(vehicles):
        raise ValueError("Missing installment schedule")

    report = {
        "schema": "Demo V1; not verified customer finance format",
        "fictional_only": True, "export_allowed": True,
        "exception_count": 0, "exceptions": [], "month": month,
        "vehicles": [{"vehicle_id": key, "registration": item["registration"],
                      "result": "MATCH", "xlsx_source": item["xlsx_source"],
                      "contract_source": contracts[key]["source"]}
                     for key, item in vehicles.items()],
        "payments": payments,
        "all_payments_total_jpy": sum(p["amount_jpy"] for p in payments),
        "selected_total_jpy": sum(p["amount_jpy"] for p in payments
                                  if p["payment_date"][:7] == month),
        "source_inputs": [{"name": p.name, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}
                          for p in inputs],
        "pdf_pages": {"contracts": len(contract_pages), "installments": len(schedule_pages)},
    }
    WRITER.validated_rows(report)
    return report


def run(vehicles, contracts, installments, month, out):
    output = Path(out).resolve()
    if output.exists():
        raise FileExistsError("Output destination already exists; inspect it before another run")
    report = extract_report(vehicles, contracts, installments, month)
    workbook_bytes = WRITER.build_xlsx(report, month)
    # Source validation and in-memory authoring finish before any output path.
    output.mkdir()
    workbook_path = output / "payment-schedule.xlsx"
    with workbook_path.open("xb") as target:
        target.write(workbook_bytes)
    saved_bytes = workbook_path.read_bytes()
    csv_bytes = WRITER.export_csv(saved_bytes, report, month)
    with (output / "finance.csv").open("xb") as target:
        target.write(csv_bytes)
    with (output / "source-report.json").open("x", encoding="utf-8") as target:
        json.dump(report, target, ensure_ascii=False, indent=2)
        target.write("\n")
    return {"status": "FICTIONAL_THREE_STEP_INTEGRATION_VERIFIED",
            "vehicles": len(report["vehicles"]), "payment_rows": len(report["payments"]),
            "month": month, "selected_total_jpy": report["selected_total_jpy"],
            "saved_xlsx_sha256": hashlib.sha256(saved_bytes).hexdigest(),
            "csv_sha256": hashlib.sha256(csv_bytes).hexdigest(),
            "fictional_only": True, "client_compatibility_verified": False}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ("vehicles", "contracts", "installments", "month", "out"):
        parser.add_argument("--" + name, required=True)
    args = parser.parse_args()
    try:
        result = run(**vars(args))
    except (ValueError, OSError, KeyError, TypeError) as error:
        print(json.dumps({"status": "BLOCKED", "error": str(error)}), file=sys.stderr)
        return 2
    print(json.dumps(result))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
