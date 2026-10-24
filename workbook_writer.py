"""Payment workbook writer for the fixed fictional demonstration schema.

The saved workbook is reread before CSV export. No real client format,
OCR, GUI or external AI service is supplied.
"""
from __future__ import annotations
import csv
import hashlib
import io
import re
from datetime import date, datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FIELDS = ["payment_date", "vehicle_id", "registration", "amount_jpy"]
HEADER = ["Vehicle ID", "Registration", "Instalment", "Due date", "Amount JPY", "PDF source"]

def integer(value):
    if type(value) is not int or value < 0:
        raise ValueError("Expected nonnegative integer JPY/count")
    return value

def text(value):
    if not isinstance(value, str) or not value.strip():
        raise ValueError("Missing string field")
    return value

def month_key(month):
    if not isinstance(month, str) or not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", month):
        raise ValueError("Expected YYYY-MM")
    date(int(month[:4]), int(month[5:]), 1)
    return month

def validated_rows(report):
    if (report.get("schema") != "Demo V1; not verified customer finance format"
            or report.get("export_allowed") is not True
            or type(report.get("exception_count")) is not int
            or report.get("exception_count") != 0):
        raise ValueError("Source exceptions block export")
    if report.get("exceptions") != [] or report.get("fictional_only") is not True:
        raise ValueError("This component accepts only the validated fictional schema")
    vehicles = {}
    for item in report["vehicles"]:
        key, reg = text(item["vehicle_id"]), text(item["registration"])
        if key in vehicles or item["result"] != "MATCH":
            raise ValueError("Duplicate/unmatched vehicle")
        vehicles[key] = reg
    if not vehicles:
        raise ValueError("No validated vehicles")
    rows = []
    seen = set()
    for item in report["payments"]:
        key, reg = text(item["vehicle_id"]), text(item["registration"])
        if key not in vehicles or vehicles[key] != reg:
            raise ValueError("Payment identity does not match vehicle source")
        number = integer(item["installment"])
        if number < 1 or (key, number) in seen:
            raise ValueError("Duplicate/invalid installment")
        seen.add((key, number))
        due_text = text(item["payment_date"])
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", due_text):
            raise ValueError("Expected ISO calendar date")
        due = date.fromisoformat(due_text)
        rows.append((key, reg, number, due, integer(item["amount_jpy"]), text(item["source"])))
    if not rows or {row[0] for row in rows} != set(vehicles):
        raise ValueError("Missing payment coverage")
    for key in vehicles:
        numbers = sorted(r[2] for r in rows if r[0] == key)
        if numbers != list(range(1, len(numbers)+1)):
            raise ValueError("Non-contiguous installment sequence")
    if sum(row[4] for row in rows) != integer(report["all_payments_total_jpy"]):
        raise ValueError("Source total does not reconcile")
    original_month = month_key(report["month"])
    selected = sum(row[4] for row in rows if row[3].strftime("%Y-%m") == original_month)
    if selected != integer(report["selected_total_jpy"]):
        raise ValueError("Original selected-month total does not reconcile")
    return sorted(rows, key=lambda row:(row[3], row[0], row[2]))

def put(sheet, row, col, value):
    cell = sheet.cell(row, col, value)
    # Source identifiers/text must remain text, including strings beginning "=".
    if isinstance(value, str):
        cell.data_type = "s"
    return cell

def style(sheet, widths, header):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = f"A{header+1}"
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    for row in sheet:
        sheet.row_dimensions[row[0].row].height = 22
        for cell in row:
            cell.font = Font(name="Arial",size=11,color="172236")
            cell.alignment = Alignment(vertical="center")
    for cell in sheet[header]:
        cell.font = Font(name="Arial",size=11,color="FFFFFF",bold=True)
        cell.fill = PatternFill("solid",fgColor="20364F")
    sheet.auto_filter.ref = f"A{header}:{get_column_letter(sheet.max_column)}{sheet.max_row}"

def build_xlsx(report, selected_month):
    """Return a new saved-result workbook; do not modify the supplied report."""
    rows = validated_rows(report)
    month = month_key(selected_month)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Payments"
    put(sheet,1,1,"Fictional payment schedule")
    for col,value in enumerate(["Selected month",month,"Monthly export JPY",
            sum(r[4] for r in rows if r[3].strftime("%Y-%m")==month),"CSV status","Allowed"],1):
        put(sheet,2,col,value)
    put(sheet,3,1,"Saved result; rerun the application to refresh. This sheet contains all months.")
    for col,value in enumerate(HEADER,1):
        put(sheet,4,col,value)
    for index,row in enumerate(rows,5):
        for col,value in enumerate(row,1):
            put(sheet,index,col,value)
        sheet.cell(index,4).number_format = "yyyy-mm-dd"
        sheet.cell(index,5).number_format = "#,##0"
    total_row = 5+len(rows)
    put(sheet,total_row,1,"All payments total")
    put(sheet,total_row,5,sum(r[4] for r in rows))
    sheet.cell(total_row,5).number_format = "#,##0"
    sheet.cell(2,4).number_format = "#,##0"
    style(sheet,[23,31,16,19,20,36],4)
    sheet["A1"].font = Font(name="Arial",size=15,bold=True,color="172236")
    sheet["A3"].font = Font(name="Arial",size=10,italic=True,color="546274")

    reconciliation = workbook.create_sheet("Reconciliation")
    values = [["Vehicle ID","Registration","Result","XLSX source","Contract source"]]
    values += [[v[k] for k in ["vehicle_id","registration","result","xlsx_source","contract_source"]]
               for v in report["vehicles"]]
    for row_num,row in enumerate(values,1):
        for col,value in enumerate(row,1):
            put(reconciliation,row_num,col,value)
    style(reconciliation,[22,31,17,42,34],1)
    exceptions = workbook.create_sheet("Exceptions")
    for row_num,row in enumerate([["Code","Vehicle ID","Source","Detail"],
                                  ["No exceptions in this fictional run","","",""]],1):
        for col,value in enumerate(row,1):
            put(exceptions,row_num,col,value)
    style(exceptions,[42,22,36,62],1)
    stream = io.BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()

def read_saved_rows(blob, report):
    """Read the actual generated XLSX, reject stale/tampered payment rows."""
    expected = validated_rows(report)
    workbook = load_workbook(io.BytesIO(blob),read_only=True,data_only=True)
    try:
        if workbook.sheetnames != ["Payments","Reconciliation","Exceptions"]:
            raise ValueError("Unexpected workbook structure")
        sheet = workbook["Payments"]
        if list(next(sheet.iter_rows(min_row=4,max_row=4,values_only=True))) != HEADER:
            raise ValueError("Unexpected payment header")
        rows, found_total = [], False
        for raw in sheet.iter_rows(min_row=5,values_only=True):
            key, reg, number, due, amount, source = raw
            if key == "All payments total":
                if integer(amount) != sum(r[4] for r in rows):
                    raise ValueError("Saved total does not reconcile")
                found_total = True
                continue
            if found_total:
                if any(value is not None for value in raw):
                    raise ValueError("Unexpected data after total")
                continue
            if isinstance(due,datetime):
                due = due.date()
            if not isinstance(due,date):
                raise ValueError("Saved date is not typed as a date")
            rows.append((text(key),text(reg),integer(number),due,integer(amount),text(source)))
        if not found_total or sorted(rows) != sorted(expected):
            raise ValueError("Saved detail differs from validated source")
        return sorted(rows,key=lambda row:(row[3],row[0],row[2]))
    finally:
        workbook.close()

def export_csv(blob, report, month):
    """Export the chosen month from actual saved detail, not JSON substitution."""
    month = month_key(month)
    rows = read_saved_rows(blob,report)
    selected = [r for r in rows if r[3].strftime("%Y-%m")==month]
    output = io.StringIO(newline="")
    writer = csv.writer(output,lineterminator="\r\n")
    writer.writerow(FIELDS)
    writer.writerows((r[3].isoformat(),r[0],r[1],r[4]) for r in selected)
    return output.getvalue().encode("utf-8-sig")
